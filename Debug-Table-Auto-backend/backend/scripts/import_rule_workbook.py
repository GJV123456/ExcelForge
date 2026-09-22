"""Import the rule workbook into the ExcelForge SQLite catalog.

Usage (from the backend directory):
    python scripts/import_rule_workbook.py "C:\\path\\操作步骤(1).xlsx"

The importer replaces catalog business rows in one transaction while keeping
the schema and generation-history tables intact.
"""

from __future__ import annotations

import argparse
import hashlib
import re
import sqlite3
from pathlib import Path

from openpyxl import load_workbook

from app.db.catalog import CatalogRepository
from app.resource_paths import resolve_database_path
from scripts.seed_equipment_hierarchy import seed as seed_equipment_hierarchy


def _text(value: object) -> str:
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip() if value is not None else ""


def _code(prefix: str, value: str) -> str:
    digest = hashlib.sha1(value.encode("utf-8")).hexdigest()[:12].upper()
    return f"{prefix}_{digest}"


def _discipline(value: str) -> str:
    return value or "自动化"


def _insert_rule(
    connection: sqlite3.Connection,
    *,
    equipment_type: str,
    content: str,
    steps: list[tuple[str, list[tuple[str, list[str]]]]],
    type_ids: dict[str, int],
    sequence: int,
) -> None:
    if not equipment_type or not steps:
        return
    type_id = type_ids.get(equipment_type)
    if type_id is None:
        cursor = connection.execute(
            "INSERT INTO equipment_types(code, name, sort_order, is_enabled) VALUES (?, ?, ?, 1)",
            (_code("TYPE", equipment_type), equipment_type, len(type_ids) + 1),
        )
        type_id = int(cursor.lastrowid)
        type_ids[equipment_type] = type_id

    content = content or equipment_type
    content_code = _code("RULE", f"{equipment_type}:{sequence}:{content}")
    cursor = connection.execute(
        """
        INSERT INTO test_contents(
            code, equipment_type_id, content, sort_order, is_enabled
        ) VALUES (?, ?, ?, ?, 1)
        """,
        (content_code, type_id, content, sequence),
    )
    content_id = int(cursor.lastrowid)
    for step_order, (description, results) in enumerate(steps, start=1):
        if not description:
            continue
        step_cursor = connection.execute(
            """
            INSERT INTO operation_steps(
                test_content_id, description, ratio_change, sort_order, is_enabled
            ) VALUES (?, ?, '否', ?, 1)
            """,
            (content_id, description, step_order),
        )
        step_id = int(step_cursor.lastrowid)
        for result_order, (result_description, disciplines) in enumerate(results, start=1):
            if not result_description:
                continue
            result_cursor = connection.execute(
                """
                INSERT INTO result_outputs(
                    operation_step_id, description, sort_order, is_enabled
                ) VALUES (?, ?, ?, 1)
                """,
                (step_id, result_description, result_order),
            )
            result_id = int(result_cursor.lastrowid)
            for discipline_order, discipline in enumerate(disciplines or ["自动化"], start=1):
                name = _discipline(discipline)
                discipline_id = connection.execute(
                    "SELECT id FROM disciplines WHERE name = ? COLLATE NOCASE", (name,)
                ).fetchone()
                if discipline_id is None:
                    discipline_cursor = connection.execute(
                        "INSERT INTO disciplines(code, name, sort_order, is_enabled) VALUES (?, ?, ?, 1)",
                        (_code("DISC", name), name, discipline_order),
                    )
                    discipline_id = (int(discipline_cursor.lastrowid),)
                connection.execute(
                    "INSERT INTO result_output_disciplines(result_output_id, discipline_id, sort_order) VALUES (?, ?, ?)",
                    (result_id, int(discipline_id[0]), discipline_order),
                )


def _parse_sheet(sheet, *, common: bool) -> list[tuple[str, str, list[tuple[str, list[tuple[str, list[str]]]]]]]:
    rules: list[tuple[str, str, list[tuple[str, list[tuple[str, list[str]]]]]]] = []
    current_type = ""
    current_content = ""
    current_steps: list[tuple[str, list[tuple[str, list[str]]]]] = []
    current_step: tuple[str, list[tuple[str, list[str]]]] | None = None

    def flush_step() -> None:
        nonlocal current_step
        if current_step is not None and current_step[0]:
            current_steps.append(current_step)
        current_step = None

    def flush_rule() -> None:
        nonlocal current_content, current_steps
        flush_step()
        if current_type and current_steps:
            rules.append((current_type, current_content, current_steps))
        current_content = ""
        current_steps = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        values = [_text(value) for value in row[:5]]
        if common:
            type_value, operation, _, output, discipline = values
            if type_value:
                if current_type and type_value != current_type:
                    flush_rule()
                current_type = type_value
            if operation:
                flush_step()
                current_step = (operation, [])
        else:
            type_value, trigger, operation, output, discipline = values
            normalized_type = type_value.split("\n", 1)[0].strip() if type_value else ""
            if normalized_type and normalized_type != current_type:
                flush_rule()
                current_type = normalized_type
            if trigger:
                if current_step is not None or current_steps:
                    flush_rule()
                current_content = trigger
            if operation:
                flush_step()
                current_step = (operation, [])
        if output:
            if current_step is None:
                current_step = ("现场测试", [])
            current_step[1].append((output, [_discipline(discipline)]))
    flush_rule()
    return rules


def import_workbook(workbook_path: Path, database_path: Path) -> tuple[int, int, int]:
    CatalogRepository(database_path).ensure_schema()
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    with sqlite3.connect(database_path) as connection:
        connection.execute("PRAGMA foreign_keys = ON")
        connection.execute("BEGIN IMMEDIATE")
        connection.execute("DELETE FROM result_output_disciplines")
        connection.execute("DELETE FROM result_outputs")
        connection.execute("DELETE FROM operation_steps")
        connection.execute("DELETE FROM test_contents")
        connection.execute("DELETE FROM equipment_subtype_statuses")
        connection.execute("DELETE FROM equipment_type_statuses")
        connection.execute("DELETE FROM equipment_type_subtypes")
        connection.execute("DELETE FROM equipment_types")
        connection.execute("DELETE FROM disciplines")
        type_ids: dict[str, int] = {}
        rule_count = step_count = result_count = 0
        sequence = 0
        for index, sheet in enumerate(workbook.worksheets):
            common = index == 0
            for equipment_type, content, steps in _parse_sheet(sheet, common=common):
                sequence += 1
                _insert_rule(
                    connection,
                    equipment_type=equipment_type,
                    content=content,
                    steps=steps,
                    type_ids=type_ids,
                    sequence=sequence,
                )
                rule_count += 1
                step_count += len(steps)
                result_count += sum(len(results) for _, results in steps)
        connection.commit()
    # Rebuild hierarchy against the freshly assigned equipment type ids.
    seed_equipment_hierarchy(database_path)
    return rule_count, step_count, result_count


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--database", type=Path)
    args = parser.parse_args()
    database_path = resolve_database_path(args.database)
    counts = import_workbook(args.workbook.resolve(), database_path)
    print(f"database={database_path}")
    print(f"rules={counts[0]} steps={counts[1]} results={counts[2]}")


if __name__ == "__main__":
    main()
